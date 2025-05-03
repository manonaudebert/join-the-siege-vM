from io import BytesIO
from werkzeug.datastructures import FileStorage
from pypdf import PdfReader
from docx import Document

import mimetypes

from PIL import Image
import pytesseract

class FileHandler:
    def extract_text(self, file):
        raise NotImplementedError()

class PdfHandler(FileHandler):
    def extract_text(self, file):
        pdf = PdfReader(BytesIO(file.read()))
        text = "\n".join(page.extract_text() for page in pdf.pages)
        return text

class ImageHandler(FileHandler):
    def extract_text(self, file):
        image = Image.open(BytesIO(file.read()))
        text = pytesseract.image_to_string(image, lang='eng')
        return text

class WordDocHandler(FileHandler):
    def extract_text(self, file):
        doc = Document(BytesIO(file.read()))
        text_sections = [para.text for para in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text_sections.append(cell.text)

        text = "\n".join(text_sections)

        return text
    
class FileHandlerFactory:

    @staticmethod
    def get_file_handler(file: FileStorage):
        filename = file.filename.lower()
        mime_type, encoding = mimetypes.guess_type(filename)
        if mime_type:
            if 'pdf' in mime_type:
                return PdfHandler()
            elif mime_type.startswith('image'):
                return ImageHandler()
            elif 'openxmlformats-officedocument' in mime_type:
                return WordDocHandler()
            
        # If mimetype doesn't recognize the type, fall back to extension. 
        fileextention = filename.split('.')[-1]
        if fileextention == 'pdf':
            return PdfHandler()
        elif fileextention in ['jpeg','jpg','png']:
            return ImageHandler()
        elif fileextention == 'docx':
            return WordDocHandler()
        else:
            raise ValueError(f'Could not get file type for {filename}')
